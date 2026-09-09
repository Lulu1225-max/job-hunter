from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace
from uuid import UUID, uuid4

import jwt
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.api.v1 import resumes as resumes_api
from app.core import security
from app.core.config import settings
from app.core.database import get_db
from app.main import app
from app.schemas.skills import DetectedSkills
from app.services.document_parser import document_parser
from app.services.resume_service import ResumeService, ResumeUploadError
from app.services import resume_service as resume_service_module


TEST_SECRET = "resume-upload-test-secret-at-least-32-bytes"
TEST_ISSUER = "https://unit-test.supabase.co/auth/v1"


def token(user_id: UUID) -> str:
    return jwt.encode(
        {
            "sub": str(user_id),
            "email": "resume@example.com",
            "aud": "authenticated",
            "iss": TEST_ISSUER,
            "role": "authenticated",
            "exp": datetime.now(timezone.utc) + timedelta(minutes=15),
        },
        TEST_SECRET,
        algorithm="HS256",
    )


class FakeDb:
    def commit(self):
        pass

    def rollback(self):
        pass


class FakeResumeApiRepo:
    def __init__(self, owner: UUID, resume_id: UUID):
        self.owner = owner
        self.resume_id = resume_id

    def list(self, db, user_id):
        return [{"id": str(self.resume_id), "user_id": str(self.owner)}] if user_id == self.owner else []

    def get(self, db, user_id, resume_id):
        if user_id == self.owner and resume_id == self.resume_id:
            return SimpleNamespace(id=resume_id, user_id=self.owner)
        return None


@pytest.fixture
def authenticated_app(monkeypatch):
    monkeypatch.setattr(settings, "supabase_url", "https://unit-test.supabase.co")
    monkeypatch.setattr(settings, "supabase_publishable_key", "test-publishable-key")
    monkeypatch.setattr(settings, "supabase_jwt_secret", TEST_SECRET)
    monkeypatch.setattr(settings, "supabase_jwt_audience", "authenticated")
    monkeypatch.setattr(
        security,
        "_decode_with_jwks",
        lambda value: (_ for _ in ()).throw(security.AuthError("No JWKS in unit test")),
    )
    monkeypatch.setattr(
        security,
        "_verify_with_auth_server",
        lambda value: (_ for _ in ()).throw(security.AuthError("No auth server in unit test")),
    )
    monkeypatch.setattr(resumes_api, "ensure_user", lambda db, user_id: None)
    app.dependency_overrides[get_db] = lambda: FakeDb()
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.clear()


def test_backend_file_validation_rejects_images_and_invalid_files(monkeypatch):
    service = ResumeService()
    monkeypatch.setattr(settings, "resume_upload_max_bytes", 8)

    with pytest.raises(ResumeUploadError, match="Only PDF and DOCX"):
        service.validate_file("portrait.png", "image/png", b"\x89PNG")
    with pytest.raises(ResumeUploadError, match="exceeds"):
        service.validate_file("resume.pdf", "application/pdf", b"%PDF-more")
    with pytest.raises(ResumeUploadError, match="not a valid PDF"):
        service.validate_file("resume.pdf", "application/pdf", b"not-pdf")
    with pytest.raises(ResumeUploadError, match="MIME"):
        service.validate_file("resume.docx", "image/jpeg", b"PK-data")


def test_docx_text_extraction_includes_paragraphs_and_tables():
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Python and SQL")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Product strategy"
    table.cell(0, 1).text = "English"
    document.save(buffer)

    text = document_parser.extract_text(buffer.getvalue(), "docx")

    assert "Python and SQL" in text
    assert "Product strategy | English" in text


def test_pdf_text_extraction_reads_page_text():
    buffer = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_reference})}
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (Python and SQL) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.write(buffer)

    text = document_parser.extract_text(buffer.getvalue(), "pdf")

    assert "Python and SQL" in text


def test_upload_uses_authenticated_user_storage_path_and_persists_results(monkeypatch):
    user_id = uuid4()
    events: list[str] = []
    rows: dict[UUID, SimpleNamespace] = {}

    class Repo:
        def create(self, db, owner_id, payload):
            events.append("create")
            row = SimpleNamespace(user_id=owner_id, **payload)
            rows[payload["id"]] = row
            return payload

        def update(self, db, owner_id, resume_id, payload):
            events.append("update")
            row = rows[resume_id]
            assert row.user_id == owner_id
            for key, value in payload.items():
                setattr(row, key, value)
            return {"id": str(resume_id), "user_id": str(owner_id), **payload}

        def get(self, db, owner_id, resume_id):
            row = rows.get(resume_id)
            return row if row and row.user_id == owner_id else None

        def delete(self, db, owner_id, resume_id):
            rows.pop(resume_id, None)
            return True

    class Storage:
        def upload(self, path, content, content_type):
            events.append("upload")
            assert path.startswith(f"{user_id}/")
            return path

        def delete(self, path):
            events.append("delete")

    monkeypatch.setattr(resume_service_module, "resumes_repo", Repo())
    monkeypatch.setattr(resume_service_module, "resume_storage", Storage())
    monkeypatch.setattr(resume_service_module.document_parser, "extract_text", lambda content, file_type: "Python SQL")
    monkeypatch.setattr(
        resume_service_module.resume_skill_extractor,
        "extract",
        lambda text: DetectedSkills(technical_skills=["Python", "SQL"]),
    )

    result = ResumeService().upload(
        FakeDb(),
        user_id,
        "resume.pdf",
        "application/pdf",
        b"%PDF-valid-test",
        None,
    )

    assert events == ["upload", "create", "update"]
    assert result["detected_skills"]["technical_skills"] == ["Python", "SQL"]
    assert result["extracted_text"] == "Python SQL"


def test_confirm_skills_requires_owned_resume_and_only_merges_selected(monkeypatch):
    user_a = uuid4()
    user_b = uuid4()
    resume_id = uuid4()
    resume = SimpleNamespace(
        user_id=user_a,
        detected_skills=DetectedSkills(
            technical_skills=["Python", "SQL"],
            tools=["Docker"],
        ).model_dump(),
    )
    stored_profile = {
        "technical_skills": ["TypeScript"],
        "product_skills": [],
        "soft_skills": [],
        "tools": [],
        "languages": [],
    }

    class Repo:
        def get(self, db, user_id, requested_id):
            return resume if user_id == user_a and requested_id == resume_id else None

    class Profiles:
        def get(self, db, user_id):
            return stored_profile

        def upsert(self, db, user_id, updates):
            stored_profile.update(updates)
            return stored_profile

    monkeypatch.setattr(resume_service_module, "resumes_repo", Repo())
    monkeypatch.setattr(resume_service_module, "profile_repo", Profiles())
    selected = DetectedSkills(technical_skills=["Python"])

    updated = ResumeService().confirm_skills(FakeDb(), user_a, resume_id, selected)

    assert updated["technical_skills"] == ["TypeScript", "Python"]
    assert updated["tools"] == []
    assert ResumeService().confirm_skills(FakeDb(), user_b, resume_id, selected) is None


def test_user_b_cannot_access_user_a_resume(authenticated_app, monkeypatch):
    user_a = uuid4()
    user_b = uuid4()
    resume_id = uuid4()
    monkeypatch.setattr(resumes_api, "resumes_repo", FakeResumeApiRepo(user_a, resume_id))
    monkeypatch.setattr(
        resumes_api,
        "serialize_model",
        lambda row: {"id": str(row.id), "user_id": str(row.user_id)},
    )

    owner_list = authenticated_app.get(
        "/api/v1/resumes",
        headers={"Authorization": f"Bearer {token(user_a)}"},
    )
    other_list = authenticated_app.get(
        "/api/v1/resumes",
        headers={"Authorization": f"Bearer {token(user_b)}"},
    )
    blocked = authenticated_app.get(
        f"/api/v1/resumes/{resume_id}",
        headers={"Authorization": f"Bearer {token(user_b)}"},
    )

    assert owner_list.status_code == 200
    assert len(owner_list.json()) == 1
    assert other_list.json() == []
    assert blocked.status_code == 404
