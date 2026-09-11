from __future__ import annotations
import csv, io, re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any
from openpyxl import load_workbook
from app.utils.matching import source_fingerprint
from app.utils.normalization import canonical_job_type, display_text, parse_date

CANONICAL_FIELDS=("company","role","location","job_url","description","deadline","job_type","industry","salary","application_start_date","campus_category","referral_available","graduation_cohort","company_type")
ALIASES={
"company":{"公司名称","公司","公司简称","企业","企业名称","Company","Company Name","Employer"},
"role":{"岗位名称","岗位","职位","职位名称","Role","Job Title","Position"},
"location":{"地区","地点","工作地点","Base地","Base","Location","City"},
"job_url":{"投递链接","职位链接","岗位链接","网申链接","Job URL","Job Link","URL","Apply URL"},
"description":{"职位画像","岗位要求","职位描述","工作内容","JD","Job Description","Description"},
"deadline":{"网申截止时间","截止时间","截止日期","Deadline","Application Deadline"},
"job_type":{"招聘类型","岗位类型","Job Type","Employment Type"},"industry":{"行业","Industry"},"salary":{"薪资","工资","Salary","Compensation"},
"application_start_date":{"网申开始时间","开始时间","网申开始日期","Application Start Date"},
"campus_category":{"校招分类","招聘分类","校招类型","Campus Category"},
"referral_available":{"是否在内推","是否内推","内推","内推状态","Referral Available"},
"graduation_cohort":{"届次","招聘届次","毕业届次","Graduation Cohort"},
"company_type":{"企业性质","公司性质","企业类型","Company Type"}}
REQUIRED_FIELDS={"company"}
def _key(v:Any)->str:return re.sub(r"[\s_\-/:：()（）]+","",(display_text(v) or "").casefold())
@dataclass(frozen=True)
class ImportIssue: row_number:int; reason:str; values:dict[str,Any]
@dataclass(frozen=True)
class NormalizedJobRow:
 source:str; company:str; role:str|None=None; location:str|None=None; job_url:str|None=None; description:str|None=None; deadline:str|None=None; job_type:str|None=None; industry:str|None=None; salary:str|None=None; application_start_date:str|None=None; campus_category:str|None=None; referral_available:bool|None=None; graduation_cohort:str|None=None; company_type:str|None=None; source_hash:str=""

def normalize_referral(value:Any)->bool|None:
 normalized=_key(value)
 if normalized in {_key(v) for v in ("是","有","可内推","内推","Yes","Y","true")}:return True
 if normalized in {_key(v) for v in ("否","无","不可内推","No","N","false")}:return False
 return None
def detect_header_row(rows:list[list[Any]])->int|None:
 keys={_key(a) for vals in ALIASES.values() for a in vals}; scores=[(i,sum(_key(c) in keys for c in row if display_text(c))) for i,row in enumerate(rows[:20])]; best=max(scores,key=lambda x:x[1],default=(0,0)); return best[0] if best[1]>=2 else None
def detect_columns(headers:list[Any])->dict[str,str]:
 lookup={_key(a):f for f,aliases in ALIASES.items() for a in aliases}; out={}; used=set()
 for h in headers:
  original=display_text(h); field=lookup.get(_key(h))
  if original and field and field not in used: out[original]=field; used.add(field)
 return out
def _read_tables(content:bytes,filename:str)->dict[str,list[list[Any]]]:
 suffix=Path(filename).suffix.lower()
 if suffix==".csv":
  try:text=content.decode("utf-8-sig",errors="strict")
  except UnicodeDecodeError as exc:raise ValueError("CSV must be UTF-8 encoded") from exc
  return {"CSV":[list(r) for r in csv.reader(io.StringIO(text))]}
 if suffix!=".xlsx":raise ValueError("Only .xlsx and .csv files are supported")
 try:
  wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True); return {ws.title:[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
 except Exception as exc:raise ValueError("The Excel workbook could not be read") from exc
def inspect_import(content:bytes,filename:str)->dict[str,Any]:
 sheets=[]
 for name,rows in _read_tables(content,filename).items():
  idx=detect_header_row(rows); headers=[] if idx is None else [display_text(v) for v in rows[idx] if display_text(v)]; mapping=detect_columns(headers)
  sheets.append({"sheet_name":name,"header_row":None if idx is None else idx+1,"headers":headers,"mapping":mapping,"confidence":round(len(set(mapping.values()))/len(CANONICAL_FIELDS),2)})
 return {"filename":Path(filename).name,"sheets":sheets}
def _validated_mapping(headers:list[str],mapping:dict[str,str|None])->dict[str,str]:
 clean={s:t for s,t in mapping.items() if t}
 if any(s not in headers for s in clean) or any(t not in CANONICAL_FIELDS for t in clean.values()):raise ValueError("Column mapping contains an unknown source or target field")
 if len(set(clean.values()))!=len(clean):raise ValueError("Each canonical field can only be mapped once")
 missing=REQUIRED_FIELDS-set(clean.values())
 if missing:raise ValueError(f"Required mapping missing: {', '.join(sorted(missing))}")
 return clean
def normalize_row(source:str,raw:dict[str,Any],row_number:int):
 company,role=display_text(raw.get("company")),display_text(raw.get("role"))
 if not company:return None,ImportIssue(row_number,"Missing company",raw)
 dv=raw.get("deadline"); deadline=parse_date(dv); sv=raw.get("application_start_date"); start_date=parse_date(sv); warning=bool((display_text(dv) and deadline is None) or (display_text(sv) and start_date is None))
 vals={f:display_text(raw.get(f)) or None for f in CANONICAL_FIELDS}; vals["deadline"]=deadline.isoformat() if deadline else None; vals["application_start_date"]=start_date.isoformat() if start_date else None; vals["job_type"]=canonical_job_type(raw.get("job_type")) or None; vals["referral_available"]=normalize_referral(raw.get("referral_available"))
 job=NormalizedJobRow(source=source,company=company,role=role,**{k:v for k,v in vals.items() if k not in {"company","role"}},source_hash=source_fingerprint([company,role,vals["location"],vals["job_url"],vals["deadline"],vals["description"]]))
 return job,(ImportIssue(row_number,"Invalid optional date; imported as blank",raw) if warning else None)
def parse_import(content:bytes,filename:str,sheet_name:str,header_row:int,mapping:dict[str,str|None])->dict[str,Any]:
 tables=_read_tables(content,filename)
 if sheet_name not in tables:raise ValueError("Selected sheet was not found")
 rows=tables[sheet_name]
 if header_row<1 or header_row>len(rows):raise ValueError("Header row is outside the file")
 headers=[display_text(v) for v in rows[header_row-1]]; mapping=_validated_mapping(headers,mapping); jobs=[]; issues=[]; warnings=[]; blank=0
 for number,values in enumerate(rows[header_row:],start=header_row+1):
  by={headers[i]:values[i] if i<len(values) else None for i in range(len(headers)) if headers[i]}; mapped={target:by.get(source) for source,target in mapping.items()}
  if not any(display_text(v) for v in mapped.values()):blank+=1;continue
  job,issue=normalize_row(Path(filename).name,mapped,number)
  if job:
   jobs.append(asdict(job))
   if issue:warnings.append(asdict(issue))
  else:issues.append(asdict(issue))
 sample=[{**j,"valid":True} for j in jobs[:10]]+[{**i["values"],"valid":False,"reason":i["reason"],"row_number":i["row_number"]} for i in issues[:10]]
 return {"total_rows":len(jobs)+len(issues)+blank,"jobs":jobs,"issues":issues,"warnings":warnings,"invalid":len(issues),"skipped_blank":blank,"sample_rows":sample[:10]}
# Path-based compatibility helpers for existing tests/scripts; API never accepts paths.
def preview_workbook(path:str|Path):
 from types import SimpleNamespace
 data=inspect_import(Path(path).read_bytes(),Path(path).name); return SimpleNamespace(filename=data["filename"],sheets=[SimpleNamespace(**s,sample_rows=[]) for s in data["sheets"]])
def import_sheet(path:str|Path,sheet_name:str,existing_hashes:set[str]|None=None):
 from types import SimpleNamespace
 data=inspect_import(Path(path).read_bytes(),Path(path).name); sheet=next(s for s in data["sheets"] if s["sheet_name"]==sheet_name); parsed=parse_import(Path(path).read_bytes(),Path(path).name,sheet_name,sheet["header_row"],sheet["mapping"]); existing=set(existing_hashes or ()); created=updated=0
 for job in parsed["jobs"]:
  if job["source_hash"] in existing:updated+=1
  else:created+=1;existing.add(job["source_hash"])
 return SimpleNamespace(total_rows=parsed["total_rows"],created=created,updated=updated,skipped=parsed["skipped_blank"],invalid=parsed["invalid"],jobs=[SimpleNamespace(**j) for j in parsed["jobs"]],issues=[SimpleNamespace(**i) for i in parsed["issues"]])
