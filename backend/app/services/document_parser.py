from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError


class DocumentParseError(ValueError):
    pass


class DocumentParser:
    def extract_text(self, content: bytes, file_type: str) -> str:
        if file_type == "pdf":
            text = self._extract_pdf(content)
        elif file_type == "docx":
            text = self._extract_docx(content)
        else:
            raise DocumentParseError("Unsupported resume file type")
        without_nuls = text.replace("\x00", "")
        normalized = "\n".join(line.rstrip() for line in without_nuls.splitlines()).strip()
        if not normalized:
            raise DocumentParseError("The document contains no extractable text")
        return normalized

    def _extract_pdf(self, content: bytes) -> str:
        try:
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise DocumentParseError("Encrypted PDF files are not supported")
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except DocumentParseError:
            raise
        except (PdfReadError, ValueError, OSError) as exc:
            raise DocumentParseError("The PDF file is invalid or unreadable") from exc

    def _extract_docx(self, content: bytes) -> str:
        try:
            document = Document(BytesIO(content))
        except (PackageNotFoundError, ValueError, KeyError) as exc:
            raise DocumentParseError("The DOCX file is invalid or unreadable") from exc

        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)


document_parser = DocumentParser()
